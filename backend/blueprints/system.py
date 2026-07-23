from __future__ import annotations

import datetime as dt
import base64
import csv
import hashlib
import io
import json
import logging
import os
import re
import shutil
import zipfile
from collections.abc import Iterable, Iterator
from decimal import Decimal, InvalidOperation
from pathlib import Path

from flask import Blueprint, jsonify, request, send_file, send_from_directory
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from auth import record_audit
from db import IS_POSTGRES

bp = Blueprint("system", __name__)
logger = logging.getLogger(__name__)
TOOLS = Path(__file__).resolve().parents[2] / "tools"


def _application():
    # O import tardio evita ciclo e mantem monkeypatches no alias legado `app`.
    from core import application

    return application


def _clear_runtime_directory(path: Path, data_root: Path) -> int:
    """Remove somente filhos de um diretorio controlado pela aplicacao."""
    resolved = path.resolve()
    resolved.relative_to(data_root.resolve())
    removed = 0
    if not resolved.exists():
        resolved.mkdir(parents=True, exist_ok=True)
        return removed
    for child in resolved.iterdir():
        if child.is_dir():
            shutil.rmtree(child)
        else:
            child.unlink()
        removed += 1
    return removed


@bp.route("/api/v1/health")
def health():
    application = _application()
    try:
        with application.db_connect() as conn:
            conn.execute("SELECT 1").fetchone()
            schema_row = conn.execute(
                "SELECT COALESCE(MAX(version), 0) FROM schema_migrations"
            ).fetchone()
    except Exception:
        logger.exception("Banco indisponivel durante health check")
        return jsonify({"status": "error", "db": "unavailable"}), 503
    return jsonify(
        {
            "status": "ok",
            "version": "1.0.0",
            "db": "postgres" if IS_POSTGRES else "sqlite",
            "commit": (os.environ.get("RENDER_GIT_COMMIT") or "")[:12],
            "schema_version": int(schema_row[0]),
            "worker_mode": application.WORKER_MODE,
            "worker_process": bool(application.IS_WORKER_PROCESS),
        }
    )


@bp.route("/ui/importar")
def import_preview_ui():
    return send_from_directory(TOOLS / "import-preview", "index.html")


@bp.route("/ui/varredura")
def import_scan_ui():
    return send_from_directory(TOOLS / "import-scan", "index.html")


@bp.route("/ui/arquivos")
def import_files_ui():
    return send_from_directory(TOOLS / "import-files", "index.html")


def _csv_bytes(columns: list[str], rows: list[tuple]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(columns)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


AUDIT_COLUMNS = [
    ("ID do lançamento", "transaction_id"),
    ("Chave técnica", "tx_key"),
    ("Data", "date"),
    ("Competência", "competence_month"),
    ("Descrição", "description"),
    ("Descrição normalizada", "description_norm"),
    ("Valor", "amount"),
    ("Tipo", "type"),
    ("Status", "status"),
    ("Conta", "account_name"),
    ("Tipo de conta", "account_type"),
    ("ID da conta", "account_id"),
    ("Livro/Centro", "ledger_name"),
    ("ID livro/centro", "ledger_id"),
    ("Categoria", "category_name"),
    ("ID categoria", "category_id"),
    ("Subcategoria", "subcategory_name"),
    ("ID subcategoria", "subcategory_id"),
    ("Parcela atual", "installment_current"),
    ("Total de parcelas", "installment_total"),
    ("Plano de parcelas", "installment_plan_id"),
    ("Valor do plano", "installment_plan_amount"),
    ("Parcelas no plano", "installment_plan_members"),
    ("Observações", "notes"),
    ("Sinalizadores", "flags"),
    ("Estabelecimento normalizado", "merchant_norm"),
    ("Método da transação", "transaction_method"),
    ("Contraparte", "counterparty_name"),
    ("Referência bancária", "bank_reference"),
    ("Bloqueado", "locked"),
    ("Classificado por", "classified_by"),
    ("Classificado em", "classified_at"),
    ("Categoria sugerida", "suggested_category_name"),
    ("ID categoria sugerida", "suggested_category_id"),
    ("Subcategoria sugerida", "suggested_subcategory_name"),
    ("ID subcategoria sugerida", "suggested_subcategory_id"),
    ("Probabilidade da sugestão", "match_probability"),
    ("Notas da sugestão", "match_notes"),
    ("ID vínculo histórico", "history_match_id"),
    ("Vínculo histórico confirmado", "history_match_confirmed"),
    ("ID histórico rejeitado", "history_match_rejected_id"),
    ("Pontuação de identidade", "identity_score"),
    ("Data histórica", "history_date"),
    ("Descrição histórica", "history_description"),
    ("Valor histórico", "history_amount"),
    ("Tipo histórico", "history_type"),
    ("Fonte histórica", "history_source_file_id"),
    ("Conta histórica", "history_account_name"),
    ("Categoria histórica", "history_category_name"),
    ("Subcategoria histórica", "history_subcategory_name"),
    ("ID conciliação", "reconciliation_id"),
    ("Papel na conciliação", "reconciliation_role"),
    ("Valor conciliado", "reconciliation_amount"),
    ("Conciliado por", "reconciliation_created_by"),
    ("Conciliado em", "reconciliation_created_at"),
    ("ID contraparte conciliada", "counterpart_id"),
    ("Data contraparte", "counterpart_date"),
    ("Descrição contraparte", "counterpart_description"),
    ("Valor contraparte", "counterpart_amount"),
    ("Tipo contraparte", "counterpart_type"),
    ("Conta contraparte", "counterpart_account_name"),
    ("ID lote de importação", "imported_file_id"),
    ("Arquivo de origem", "source_filename"),
    ("Tipo do arquivo", "source_file_type"),
    ("Hash SHA-1 do arquivo", "source_file_hash"),
    ("Banco detectado", "source_bank"),
    ("Ano do arquivo", "source_year"),
    ("Mês do arquivo", "source_month"),
    ("Natureza da fonte", "source_kind"),
    ("Importado em", "imported_at"),
    ("Linhas interpretadas", "source_total_parsed"),
    ("Linhas inseridas", "source_total_inserted"),
    ("Duplicadas no lote", "source_total_duplicates"),
    ("Erros no lote", "source_total_errors"),
    ("ID documento no cofre", "stored_document_id"),
    ("Documento no cofre", "stored_document_filename"),
    ("Competência no cofre", "stored_document_year_month"),
    ("Tamanho do documento (bytes)", "stored_document_size"),
    ("Documento guardado em", "stored_document_created_at"),
]

AUDIT_VISIBLE_COLUMNS = [
    ("Data", "date", 12),
    ("Competência", "competence_month", 13),
    ("Descrição", "description", 38),
    ("Valor", "amount", 15),
    ("Tipo", "type", 14),
    ("Status", "status", 16),
    ("Conta", "account_name", 22),
    ("Tipo de conta", "account_type", 18),
    ("Método", "transaction_method", 18),
    ("Categoria", "category_name", 22),
    ("Subcategoria", "subcategory_name", 24),
    ("Livro/Centro", "ledger_name", 20),
    ("Parcela", "installment_label", 14),
    ("Valor do plano", "installment_plan_amount", 16),
    ("Sinalizadores", "flags", 26),
    ("Observações", "notes", 32),
    ("Bloqueado", "locked", 12),
    ("Classificado por", "classified_by", 18),
    ("Classificado em", "classified_at", 20),
    ("Categoria sugerida", "suggested_category_name", 22),
    ("Subcategoria sugerida", "suggested_subcategory_name", 24),
    ("Probabilidade da sugestão", "match_probability", 18),
    ("Notas da sugestão", "match_notes", 34),
    ("Vínculo histórico confirmado", "history_match_confirmed", 18),
    ("Pontuação do vínculo", "identity_score", 18),
    ("Data histórica", "history_date", 14),
    ("Descrição histórica", "history_description", 34),
    ("Valor histórico", "history_amount", 16),
    ("Tipo histórico", "history_type", 14),
    ("Origem do histórico", "history_source_display", 22),
    ("Conta histórica", "history_account_name", 22),
    ("Categoria histórica", "history_category_name", 22),
    ("Subcategoria histórica", "history_subcategory_name", 24),
    ("Papel na conciliação", "reconciliation_role", 18),
    ("Valor conciliado", "reconciliation_amount", 16),
    ("Conciliado por", "reconciliation_created_by", 18),
    ("Conciliado em", "reconciliation_created_at", 20),
    ("Data contraparte", "counterpart_date", 14),
    ("Descrição contraparte", "counterpart_description", 34),
    ("Valor contraparte", "counterpart_amount", 16),
    ("Tipo contraparte", "counterpart_type", 14),
    ("Conta contraparte", "counterpart_account_name", 22),
    ("Arquivo de origem", "source_filename", 36),
    ("Natureza da fonte", "source_kind", 18),
    ("Banco detectado", "source_bank", 18),
    ("Importado em", "imported_at", 20),
    ("Documento no cofre", "stored_document_filename", 36),
]

AUDIT_TECHNICAL_COLUMNS = [
    ("ID do lançamento", "transaction_id", 18),
    ("Chave técnica", "tx_key", 44),
    ("Descrição normalizada", "description_norm", 32),
    ("Estabelecimento normalizado", "merchant_norm", 30),
    ("Contraparte (texto bruto)", "counterparty_name", 30),
    ("Referência bancária", "bank_reference", 24),
    ("Parcela atual", "installment_current", 14),
    ("Total de parcelas", "installment_total", 14),
    ("ID da conta", "account_id", 18),
    ("ID livro/centro", "ledger_id", 18),
    ("ID categoria", "category_id", 18),
    ("ID subcategoria", "subcategory_id", 18),
    ("ID plano de parcelas", "installment_plan_id", 18),
    ("Parcelas no plano", "installment_plan_members", 16),
    ("ID categoria sugerida", "suggested_category_id", 18),
    ("ID subcategoria sugerida", "suggested_subcategory_id", 18),
    ("ID vínculo histórico", "history_match_id", 18),
    ("ID histórico rejeitado", "history_match_rejected_id", 18),
    ("Fonte histórica (bruta)", "history_source_file_id", 30),
    ("ID conciliação", "reconciliation_id", 18),
    ("ID contraparte conciliada", "counterpart_id", 18),
    ("ID lote de importação", "imported_file_id", 18),
    ("Tipo do arquivo", "source_file_type", 16),
    ("Hash SHA-1 do arquivo", "source_file_hash", 42),
    ("Ano do arquivo", "source_year", 14),
    ("Mês do arquivo", "source_month", 14),
    ("Linhas interpretadas", "source_total_parsed", 16),
    ("Linhas inseridas", "source_total_inserted", 16),
    ("Duplicadas no lote", "source_total_duplicates", 16),
    ("Erros no lote", "source_total_errors", 14),
    ("ID documento no cofre", "stored_document_id", 18),
    ("Competência no cofre", "stored_document_year_month", 18),
    ("Tamanho do documento (bytes)", "stored_document_size", 18),
    ("Documento guardado em", "stored_document_created_at", 20),
]

AUDIT_EXPORT_COLUMNS = AUDIT_VISIBLE_COLUMNS + AUDIT_TECHNICAL_COLUMNS

AUDIT_CURRENCY_KEYS = {
    "amount",
    "installment_plan_amount",
    "history_amount",
    "reconciliation_amount",
    "counterpart_amount",
}
AUDIT_DATE_KEYS = {"date", "history_date", "counterpart_date"}
AUDIT_DATETIME_KEYS = {
    "classified_at",
    "reconciliation_created_at",
    "imported_at",
    "stored_document_created_at",
}
AUDIT_PERCENT_KEYS = {"match_probability", "identity_score"}

AUDIT_TRANSLATIONS = {
    "type": {"income": "Receita", "expense": "Despesa"},
    "status": {"pending": "Pendente", "reconciled": "Classificado"},
    "account_type": {
        "credit_card": "Cartão de crédito",
        "checking": "Conta corrente",
        "savings": "Poupança",
        "investment": "Investimento",
        "cash": "Dinheiro",
    },
    "transaction_method": {
        "credit_card": "Cartão de crédito",
        "debit_card": "Cartão de débito",
        "pix": "Pix",
        "ted": "TED",
        "boleto": "Boleto",
        "bank_transfer": "Transferência",
        "automatic_debit": "Débito automático",
        "fee": "Tarifa",
        "refund": "Estorno",
        "other": "Outro",
    },
    "source_kind": {
        "cartao": "Fatura de cartão",
        "extrato": "Extrato bancário",
    },
    "reconciliation_role": {
        "income": "Receita",
        "receita": "Receita",
        "expense": "Despesa",
        "despesa": "Despesa",
    },
}

AUDIT_FLAG_TRANSLATIONS = {
    "installment": "Parcelado",
    "inter_account": "Entre contas",
    "non_count": "Não contabilizado",
    "tax": "Imposto/Taxa",
    "cartao_debito": "Cartão de débito",
    "rendimento": "Rendimento",
    "fatura": "Pagamento de fatura",
    "card_payment": "Pagamento de fatura",
    "investment": "Investimento",
    "cashback": "Cashback",
}

# XML 1.0 nao permite estes controles dentro de uma celula. Eles podem chegar
# de textos extraidos de PDF e fariam o openpyxl abortar toda a exportacao.
_ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_AUDIT_CURRENCY_COLUMN_PATTERN = "|".join(
    get_column_letter(index)
    for index, (_, key, _) in enumerate(AUDIT_EXPORT_COLUMNS, start=1)
    if key in AUDIT_CURRENCY_KEYS
).encode("ascii")
_AUDIT_CURRENCY_XML = re.compile(
    rb'(<c r="(?:'
    + _AUDIT_CURRENCY_COLUMN_PATTERN
    + rb')\d+"[^>]*><v>)(-?\d+(?:\.\d+)?)(</v></c>)'
)


def _audit_transaction_rows(conn) -> Iterator[dict]:
    cursor = conn.execute(
        """
        SELECT
          t.id AS transaction_id,t.tx_key,t.date,t.competence_month,t.description,
          t.description_norm,t.amount,t.type,t.status,t.account_id,a.name AS account_name,
          a.type AS account_type,t.ledger_id,l.name AS ledger_name,t.category_id,
          c.name AS category_name,t.subcategory_id,s.name AS subcategory_name,
          t.installment_current,t.installment_total,t.installment_plan_id,
          ip.installment_amount AS installment_plan_amount,
          (SELECT COUNT(1) FROM transactions tp WHERE tp.installment_plan_id=t.installment_plan_id)
            AS installment_plan_members,
          t.notes,t.flags,t.merchant_norm,t.transaction_method,t.counterparty_name,
          t.bank_reference,t.locked,t.classified_by,t.classified_at,
          t.suggested_category_id,sc.name AS suggested_category_name,
          t.suggested_subcategory_id,ss.name AS suggested_subcategory_name,
          t.match_probability,t.match_notes,t.history_match_id,t.history_match_confirmed,
          t.history_match_rejected_id,t.identity_score,hm.date AS history_date,
          hm.description AS history_description,hm.amount AS history_amount,
          hm.type AS history_type,hm.source_file_id AS history_source_file_id,
          ha.name AS history_account_name,hc.name AS history_category_name,
          hs.name AS history_subcategory_name,tr.id AS reconciliation_id,
          CASE WHEN tr.expense_transaction_id=t.id THEN 'despesa' WHEN tr.income_transaction_id=t.id THEN 'receita' ELSE '' END
            AS reconciliation_role,
          tr.amount AS reconciliation_amount,tr.created_by AS reconciliation_created_by,
          tr.created_at AS reconciliation_created_at,rt.id AS counterpart_id,
          rt.date AS counterpart_date,rt.description AS counterpart_description,
          rt.amount AS counterpart_amount,rt.type AS counterpart_type,
          ra.name AS counterpart_account_name,t.imported_file_id,
          f.filename AS source_filename,f.file_type AS source_file_type,
          f.file_hash AS source_file_hash,f.bank AS source_bank,f.year AS source_year,
          f.month AS source_month,f.source_kind,f.imported_at,
          f.total_parsed AS source_total_parsed,f.total_inserted AS source_total_inserted,
          f.total_duplicates AS source_total_duplicates,f.total_errors AS source_total_errors,
          sd.id AS stored_document_id,sd.filename AS stored_document_filename,
          sd.year_month AS stored_document_year_month,sd.size AS stored_document_size,
          sd.created_at AS stored_document_created_at
        FROM transactions t
        JOIN accounts a ON a.id=t.account_id
        LEFT JOIN ledgers l ON l.id=t.ledger_id
        LEFT JOIN categories c ON c.id=t.category_id
        LEFT JOIN subcategories s ON s.id=t.subcategory_id
        LEFT JOIN categories sc ON sc.id=t.suggested_category_id
        LEFT JOIN subcategories ss ON ss.id=t.suggested_subcategory_id
        LEFT JOIN installment_plans ip ON ip.id=t.installment_plan_id
        LEFT JOIN classification_history hm ON hm.id=t.history_match_id
        LEFT JOIN accounts ha ON ha.id=hm.account_id
        LEFT JOIN categories hc ON hc.id=hm.category_id
        LEFT JOIN subcategories hs ON hs.id=hm.subcategory_id
        LEFT JOIN transaction_reconciliations tr
          ON tr.expense_transaction_id=t.id OR tr.income_transaction_id=t.id
        LEFT JOIN transactions rt
          ON rt.id=CASE WHEN tr.expense_transaction_id=t.id THEN tr.income_transaction_id
                        ELSE tr.expense_transaction_id END
        LEFT JOIN accounts ra ON ra.id=rt.account_id
        LEFT JOIN imported_files f ON f.id=t.imported_file_id
        LEFT JOIN stored_documents sd ON sd.imported_file_id=f.id
        ORDER BY t.date,t.id
        """
    )
    columns = [item[0] for item in cursor.description]
    seen: set[str] = set()
    while batch := cursor.fetchmany(500):
        for raw in batch:
            row = dict(zip(columns, tuple(raw)))
            if row["transaction_id"] in seen:
                continue
            seen.add(row["transaction_id"])
            yield row


def _excel_audit_value(key: str, value):
    if value in (None, ""):
        return ""
    if key in {"locked", "history_match_confirmed"}:
        return "Sim" if value else "Não"
    if key in AUDIT_CURRENCY_KEYS:
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except (InvalidOperation, TypeError, ValueError):
            pass
    if key in AUDIT_PERCENT_KEYS:
        try:
            return round(float(value), 2)
        except (TypeError, ValueError):
            pass
    if key in AUDIT_DATE_KEYS:
        if isinstance(value, dt.datetime):
            return value.date()
        if isinstance(value, str):
            try:
                return dt.date.fromisoformat(value[:10])
            except ValueError:
                pass
    if key in AUDIT_DATETIME_KEYS:
        if isinstance(value, str):
            try:
                value = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
            except ValueError:
                pass
        if isinstance(value, dt.datetime) and value.tzinfo is not None:
            return value.astimezone(dt.timezone.utc).replace(tzinfo=None)
    if isinstance(value, str):
        return _ILLEGAL_EXCEL_CHARACTERS.sub(" ", value)
    return value


def _audit_history_source(value) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    if ":sheet:entradas" in lowered:
        return "Planilha Entradas"
    if ":sheet:saidas" in lowered:
        return "Planilha Saídas"
    if lowered.startswith("manual:"):
        return "Lançamento manual"
    return text


def _audit_flags(value) -> str:
    translated = []
    for item in re.split(r"[,;]", str(value or "")):
        clean = item.strip().strip("[]'\"")
        if not clean:
            continue
        translated.append(AUDIT_FLAG_TRANSLATIONS.get(clean.lower(), clean))
    return "; ".join(translated)


def _audit_display_value(key: str, row: dict):
    if key == "installment_label":
        current = row.get("installment_current")
        total = row.get("installment_total")
        value = f"{current} de {total}" if current and total else ""
    elif key == "history_source_display":
        value = _audit_history_source(row.get("history_source_file_id"))
    else:
        value = row.get(key)

    if key == "flags":
        value = _audit_flags(value)
    translation_key = (
        "type" if key in {"history_type", "counterpart_type"} else key
    )
    translations = AUDIT_TRANSLATIONS.get(translation_key)
    if translations and value not in (None, ""):
        value = translations.get(str(value).strip().lower(), value)
    return _excel_audit_value(key, value)


def _normalize_audit_currency_xml(output: io.BytesIO) -> io.BytesIO:
    """Mantem valores monetarios numericos e exatos em qualquer leitor XLSX."""

    def replace(match: re.Match[bytes]) -> bytes:
        value = Decimal(match.group(2).decode("ascii")).quantize(Decimal("0.01"))
        return match.group(1) + format(value, "f").encode("ascii") + match.group(3)

    normalized = io.BytesIO()
    output.seek(0)
    with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            if item.filename == "xl/worksheets/sheet1.xml":
                target.writestr(
                    item,
                    _AUDIT_CURRENCY_XML.sub(replace, source.read(item)),
                )
                continue
            with (
                source.open(item) as source_file,
                target.open(item, "w") as target_file,
            ):
                shutil.copyfileobj(source_file, target_file)
    output.close()
    normalized.seek(0)
    return normalized


def _audit_workbook(rows: Iterable[dict], created_at: str) -> io.BytesIO:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Lançamentos")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.tabColor = "2563EB"
    headers = [label for label, _, _ in AUDIT_EXPORT_COLUMNS]
    dark = "172033"
    technical_dark = "3A4358"
    light = "DCE6F1"
    lighter = "F6F8FC"
    banded = "F8FAFC"
    white = "FFFFFF"
    thin = Side(style="thin", color="D7DCE5")
    banded_fill = PatternFill("solid", fgColor=banded)
    header = []
    for index, label in enumerate(headers):
        cell = WriteOnlyCell(sheet, value=label)
        cell.fill = PatternFill(
            "solid",
            fgColor=dark if index < len(AUDIT_VISIBLE_COLUMNS) else technical_dark,
        )
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(bottom=thin)
        header.append(cell)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "D2"

    for index, (_, _, width) in enumerate(AUDIT_EXPORT_COLUMNS, start=1):
        dimension = sheet.column_dimensions[get_column_letter(index)]
        dimension.width = width
        if index > len(AUDIT_VISIBLE_COLUMNS):
            dimension.hidden = True
            dimension.outlineLevel = 1
    sheet.sheet_properties.outlinePr.summaryRight = True
    sheet.append(header)

    summary_values = {
        "total": 0,
        "income": 0.0,
        "expense": 0.0,
        "pending": 0,
        "classified": 0,
        "category": 0,
        "history": 0,
        "reconciled": 0,
        "document": 0,
    }
    for row in rows:
        summary_values["total"] += 1
        amount = float(row.get("amount") or 0)
        if row.get("type") == "income":
            summary_values["income"] += abs(amount)
        elif row.get("type") == "expense":
            summary_values["expense"] -= abs(amount)
        is_pending = row.get("status") == "pending"
        is_classified = row.get("status") == "reconciled"
        summary_values["pending"] += is_pending
        summary_values["classified"] += is_classified
        summary_values["category"] += bool(row.get("category_id"))
        summary_values["history"] += bool(row.get("history_match_id"))
        summary_values["reconciled"] += bool(row.get("reconciliation_id"))
        summary_values["document"] += bool(row.get("stored_document_id"))

        values = []
        for column_index, (_, key, _) in enumerate(AUDIT_EXPORT_COLUMNS):
            cell = WriteOnlyCell(sheet, value=_audit_display_value(key, row))
            if key in AUDIT_CURRENCY_KEYS:
                cell.number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
            elif key in AUDIT_DATE_KEYS:
                cell.number_format = "dd/mm/yyyy"
            elif key in AUDIT_DATETIME_KEYS:
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            elif key in AUDIT_PERCENT_KEYS:
                cell.number_format = '0.0"%"'
            if (
                summary_values["total"] % 2 == 0
                and column_index < len(AUDIT_VISIBLE_COLUMNS)
            ):
                cell.fill = banded_fill
            values.append(cell)
        sheet.append(values)

    final_row = int(summary_values["total"]) + 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{final_row}"

    status_col = headers.index("Status") + 1
    status_letter = get_column_letter(status_col)
    if summary_values["total"]:
        status_range = f"{status_letter}2:{status_letter}{final_row}"
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'{status_letter}2="Pendente"'],
                fill=PatternFill("solid", fgColor="FFF2CC"),
            ),
        )
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'{status_letter}2="Classificado"'],
                fill=PatternFill("solid", fgColor="E2F0D9"),
            ),
        )
        value_letter = get_column_letter(headers.index("Valor") + 1)
        value_range = f"{value_letter}2:{value_letter}{final_row}"
        sheet.conditional_formatting.add(
            value_range,
            FormulaRule(
                formula=[f"{value_letter}2<0"],
                font=Font(color="C62828"),
            ),
        )
        sheet.conditional_formatting.add(
            value_range,
            FormulaRule(
                formula=[f"{value_letter}2>0"],
                font=Font(color="16803A"),
            ),
        )

    summary = workbook.create_sheet("Resumo")
    summary.sheet_view.showGridLines = False
    summary.sheet_view.zoomScale = 100
    summary.sheet_properties.tabColor = "16A34A"
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 30
    summary.column_dimensions["C"].width = 34
    summary.column_dimensions["D"].width = 20
    total = int(summary_values["total"])
    income_raw = round(float(summary_values["income"]), 2)
    expense_raw = round(float(summary_values["expense"]), 2)
    net_raw = round(income_raw + expense_raw, 2)
    income = _excel_audit_value("amount", income_raw)
    expense = _excel_audit_value("amount", expense_raw)
    net = _excel_audit_value("amount", net_raw)
    ratio = lambda value: (float(value) / total) if total else 0.0
    generated_at = _excel_audit_value("imported_at", created_at)
    summary_rows = [
        ("Auditoria de lançamentos", "", "", ""),
        ("Gerado em (UTC)", generated_at, "Escopo", "Snapshot completo"),
        ("Total de lançamentos", total, "Classificados", summary_values["classified"]),
        ("Pendentes", summary_values["pending"], "Cobertura de classificação", ratio(summary_values["classified"])),
        ("Total de receitas", income, "Com categoria", ratio(summary_values["category"])),
        ("Total de despesas", expense, "Com vínculo histórico", ratio(summary_values["history"])),
        (
            "Saldo líquido",
            net,
            "Original preservado no cofre",
            ratio(summary_values["document"]),
        ),
        ("Conciliados", summary_values["reconciled"], "Cobertura de conciliação", ratio(summary_values["reconciled"])),
        (
            "Como usar",
            "Filtre a aba Lançamentos. As colunas técnicas ficam ocultas à direita e podem ser reexibidas.",
            "",
            "",
        ),
    ]
    for index, row_values in enumerate(summary_rows, start=1):
        cells = [WriteOnlyCell(summary, value=value) for value in row_values]
        if index == 1:
            for cell in cells:
                cell.fill = PatternFill("solid", fgColor=dark)
                cell.font = Font(color=white, bold=True, size=14)
            summary.row_dimensions[1].height = 26
        else:
            for label_cell in (cells[0], cells[2]):
                label_cell.font = Font(bold=True)
                label_cell.fill = PatternFill("solid", fgColor=light)
            cells[1].fill = PatternFill("solid", fgColor=lighter)
            cells[3].fill = PatternFill("solid", fgColor=lighter)
        if index == 2:
            cells[1].number_format = "dd/mm/yyyy hh:mm:ss"
        if index in {5, 6, 7}:
            cells[1].number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
        if index in {4, 5, 6, 7, 8}:
            cells[3].number_format = "0.0%"
        if index == 9:
            cells[0].alignment = Alignment(vertical="top")
            cells[1].alignment = Alignment(wrap_text=True, vertical="top")
            summary.row_dimensions[9].height = 42
        summary.append(cells)

    legend = workbook.create_sheet("Legenda")
    legend.sheet_view.showGridLines = False
    legend.sheet_properties.tabColor = "F59E0B"
    legend.freeze_panes = "A2"
    legend.column_dimensions["A"].width = 28
    legend.column_dimensions["B"].width = 26
    legend.column_dimensions["C"].width = 24
    legend_header = []
    for value in ("Valor exibido", "Valor interno", "Campo"):
        cell = WriteOnlyCell(legend, value=value)
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color=white, bold=True)
        legend_header.append(cell)
    legend.append(legend_header)
    legend_rows = [
        ("Despesa", "expense", "Tipo"),
        ("Receita", "income", "Tipo"),
        ("Pendente", "pending", "Status"),
        ("Classificado", "reconciled", "Status"),
        ("Cartão de crédito", "credit_card", "Tipo de conta / método"),
        ("Conta corrente", "checking", "Tipo de conta"),
        ("Pix", "pix", "Método"),
        ("TED", "ted", "Método"),
        ("Boleto", "boleto", "Método"),
        ("Transferência", "bank_transfer", "Método"),
        ("Débito automático", "automatic_debit", "Método"),
        ("Tarifa", "fee", "Método"),
        ("Estorno", "refund", "Método"),
        ("Parcelado", "INSTALLMENT", "Sinalizadores"),
        ("Entre contas", "INTER_ACCOUNT", "Sinalizadores"),
        ("Não contabilizado", "non_count", "Sinalizadores"),
        ("Imposto/Taxa", "tax", "Sinalizadores"),
        ("Rendimento", "rendimento", "Sinalizadores"),
        ("Pagamento de fatura", "fatura", "Sinalizadores"),
        ("Fatura de cartão", "cartao", "Natureza da fonte"),
        ("Extrato bancário", "extrato", "Natureza da fonte"),
    ]
    for index, row_values in enumerate(legend_rows, start=2):
        cells = [WriteOnlyCell(legend, value=value) for value in row_values]
        if index % 2 == 0:
            for cell in cells:
                cell.fill = banded_fill
        legend.append(cells)
    legend.auto_filter.ref = f"A1:C{len(legend_rows) + 1}"

    dictionary = workbook.create_sheet("Dicionário")
    dictionary.sheet_view.showGridLines = False
    dictionary.sheet_properties.tabColor = "64748B"
    dictionary.freeze_panes = "A2"
    dictionary.column_dimensions["A"].width = 38
    dictionary.column_dimensions["B"].width = 34
    dictionary.column_dimensions["C"].width = 24
    dictionary.column_dimensions["D"].width = 44
    dictionary_header = []
    for value in ("Coluna", "Campo técnico", "Bloco", "Observação"):
        cell = WriteOnlyCell(dictionary, value=value)
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color=white, bold=True)
        dictionary_header.append(cell)
    dictionary.append(dictionary_header)
    translated_keys = set(AUDIT_TRANSLATIONS) | {
        "flags",
        "installment_label",
        "history_source_display",
        "history_type",
        "counterpart_type",
    }
    for index, (label, key, _) in enumerate(AUDIT_EXPORT_COLUMNS, start=2):
        visible = index - 2 < len(AUDIT_VISIBLE_COLUMNS)
        source_key = {
            "installment_label": "installment_current + installment_total",
            "history_source_display": "history_source_file_id",
        }.get(key, key)
        if visible:
            observation = (
                "Valor traduzido/derivado para leitura."
                if key in translated_keys
                else ""
            )
            block = "Negócio (visível)"
        else:
            observation = "Oculta por padrão; reexiba o grupo para consultar."
            block = "Técnica (oculta)"
        cells = [
            WriteOnlyCell(dictionary, value=label),
            WriteOnlyCell(dictionary, value=source_key),
            WriteOnlyCell(dictionary, value=block),
            WriteOnlyCell(dictionary, value=observation),
        ]
        if index % 2 == 0:
            for cell in cells:
                cell.fill = banded_fill
        dictionary.append(cells)
    dictionary.auto_filter.ref = f"A1:D{len(AUDIT_EXPORT_COLUMNS) + 1}"
    workbook.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    workbook.save(output)
    return _normalize_audit_currency_xml(output)


@bp.route("/api/v1/system/audit-export")
def system_audit_export():
    """Exporta uma linha por lancamento com todos os vinculos de auditoria."""
    application = _application()
    forbidden = application.require_admin()
    if forbidden:
        return forbidden
    created_at = dt.datetime.now(dt.timezone.utc).isoformat()
    try:
        with application.db_connect() as conn:
            if IS_POSTGRES:
                conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
            workbook = _audit_workbook(_audit_transaction_rows(conn), created_at)
    except Exception:
        logger.exception("Falha ao gerar planilha de auditoria")
        return jsonify(
            {
                "detail": "Não foi possível gerar a planilha de auditoria",
                "code": "AUDIT_EXPORT_FAILED",
            }
        ), 500
    filename = f"auditoria-lancamentos-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return send_file(
        workbook,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )


@bp.route("/api/v1/system/audit-snapshot")
def system_audit_snapshot():
    """Exporta somente dados necessarios para auditoria financeira."""
    application = _application()
    forbidden = application.require_admin()
    if forbidden:
        return forbidden

    table_names = (
        "accounts",
        "account_file_coverage",
        "imported_files",
        "maintenance_log",
        "transaction_reconciliations",
        "transactions",
    )
    exported: dict[str, tuple[list[str], list[tuple]]] = {}
    stored_rows: list[tuple] = []
    stored_columns = [
        "id", "account_name", "year_month", "filename", "size",
        "content_sha1", "created_at", "imported_file_id",
    ]
    with application.db_connect() as conn:
        if IS_POSTGRES:
            conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
        for table in table_names:
            cursor = conn.execute(f"SELECT * FROM {table}")
            columns = [item[0] for item in cursor.description]
            exported[table] = (columns, [tuple(row) for row in cursor.fetchall()])
        docs = conn.execute(
            """
            SELECT id,account_name,year_month,filename,size,content_b64,
                   created_at,imported_file_id
            FROM stored_documents
            """
        ).fetchall()
        for row in docs:
            try:
                content_hash = hashlib.sha1(base64.b64decode(row[5])).hexdigest()
            except Exception:
                content_hash = "INVALID_BASE64"
            stored_rows.append(
                (row[0], row[1], row[2], row[3], row[4], content_hash,
                 row[6], row[7])
            )

    created_at = dt.datetime.now(dt.timezone.utc).isoformat()
    archive = io.BytesIO()
    manifest = {
        "format": "conciliador-audit-export-v1",
        "created_at": created_at,
        "database": "postgres" if IS_POSTGRES else "sqlite",
        "tables": {name: len(rows) for name, (_, rows) in exported.items()},
    }
    manifest["tables"]["stored_documents"] = len(stored_rows)
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr("manifest.json", json.dumps(manifest, indent=2, ensure_ascii=False))
        for name, (columns, rows) in exported.items():
            bundle.writestr(f"tables/{name}.csv", _csv_bytes(columns, rows))
        bundle.writestr(
            "tables/stored_documents.csv",
            _csv_bytes(stored_columns, stored_rows),
        )
    archive.seek(0)
    filename = f"production-audit-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.zip"
    return send_file(
        archive,
        mimetype="application/zip",
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )


@bp.route("/api/v1/system/reset", methods=["POST"])
def system_reset():
    application = _application()
    data = request.get_json(force=True) or {}
    if data.get("confirm") != "RESETAR":
        return jsonify(
            {"detail": "Confirmacao obrigatoria: RESETAR", "code": "VALIDATION_ERROR"}
        ), 400

    if not IS_POSTGRES:
        backups = application.DATA / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        backup = backups / (
            "conciliador_pro_before_system_reset_"
            f"{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.db"
        )
        shutil.copy2(application.DB, backup)
        backup_ref = application.project_relative_path(backup)
    else:
        backup_ref = "hosted: use o backup/branch do provedor Postgres antes de resetar"

    scope = (data.get("scope") or "transactions").strip()
    wipe_categories = bool(data.get("wipe_categories"))
    if scope not in ("transactions", "all"):
        return jsonify(
            {"detail": "scope deve ser transactions ou all", "code": "VALIDATION_ERROR"}
        ), 400

    with application.db_connect() as conn:
        before = {
            table: conn.execute(f"SELECT COUNT(1) FROM {table}").fetchone()[0]
            for table in (
                "transactions",
                "import_previews",
                "import_jobs",
                "imported_files",
                "stored_documents",
                "classification_history",
            )
        }
        for table in (
            "transaction_suggestions",
            "transaction_suggestion_state",
            "suggestion_jobs",
            "transaction_reconciliations",
            "transactions",
            "installment_plans",
            "import_jobs",
            "import_previews",
            "stored_documents",
            "account_file_coverage",
        ):
            conn.execute(f"DELETE FROM {table}")
        if scope == "all":
            conn.execute("DELETE FROM classification_history")
            conn.execute("DELETE FROM imported_files")
            if wipe_categories:
                conn.execute("DELETE FROM subcategories")
                conn.execute("DELETE FROM categories")
        else:
            conn.execute(
                "DELETE FROM imported_files WHERE file_type NOT IN ('xlsx-seed','pdf-seed')"
            )
        record_audit(
            application.current_user(),
            "system_reset",
            "system",
            detail=f"scope={scope}, wipe_categories={wipe_categories}, antes={before}",
            conn=conn,
        )
    local_files_removed = 0
    for runtime_dir in (application.UPLOADS, application.DOCS):
        try:
            local_files_removed += _clear_runtime_directory(
                runtime_dir, application.DATA
            )
        except OSError:
            logger.warning(
                "Falha ao limpar diretorio local apos reset: %s",
                runtime_dir,
                exc_info=True,
            )
    if scope == "all" and wipe_categories:
        application.seed()
    return jsonify(
        {
            "ok": True,
            "backup": backup_ref,
            "scope": scope,
            "wipe_categories": wipe_categories,
            "deleted": before,
            "local_entries_removed": local_files_removed,
        }
    )
