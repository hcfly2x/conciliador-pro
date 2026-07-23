from __future__ import annotations

import datetime as dt
import io
import os
import subprocess
import sys
import time
import unittest
import uuid
from pathlib import Path


BACKEND = Path(__file__).resolve().parents[1]
DATABASE_URL = (os.environ.get("DATABASE_URL") or "").strip()


@unittest.skipUnless(
    DATABASE_URL.startswith(("postgres://", "postgresql://")),
    "requer PostgreSQL descartavel configurado em DATABASE_URL",
)
class PostgresWorkerProcessIntegrationTests(unittest.TestCase):
    def test_audit_export_generates_xlsx_with_postgres_cursor(self) -> None:
        from openpyxl import load_workbook

        import app as app_module

        client = app_module.app.test_client()
        response = client.get("/api/v1/system/audit-export")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Lançamentos", "Resumo", "Dicionário"])
        self.assertEqual(workbook["Resumo"]["B3"].value, 0)

    def test_worker_survives_web_restart_with_persisted_job(self) -> None:
        import psycopg

        env = os.environ.copy()
        env.update(
            {
                "AUTH_DISABLED": "1",
                "WORKER_MODE": "process",
                "WORKER_POLL_SECONDS": "0.2",
            }
        )
        web = subprocess.Popen(
            [
                sys.executable,
                "-c",
                "import app, time; print('web-ready', flush=True); time.sleep(60)",
            ],
            cwd=BACKEND,
            env=env,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        job_id = str(uuid.uuid4())
        try:
            ready = False
            for _ in range(20):
                if web.stdout.readline().strip() == "web-ready":
                    ready = True
                    break
            self.assertTrue(ready, "processo web nao concluiu a inicializacao")
            with psycopg.connect(DATABASE_URL) as conn:
                conn.execute(
                    """
                    INSERT INTO recalculation_jobs(
                      id,status,processed,total,updated,error,created_at,started_at,finished_at
                    ) VALUES (%s,'running',0,0,0,'',%s,%s,'')
                    """,
                    (
                        job_id,
                        dt.datetime.now(dt.timezone.utc).isoformat(),
                        dt.datetime.now(dt.timezone.utc).isoformat(),
                    ),
                )

            worker_env = {**env, "WORKER_PROCESS": "1", "WORKER_RUN_ONCE": "1"}
            worker = subprocess.Popen(
                [sys.executable, "worker.py"],
                cwd=BACKEND,
                env=worker_env,
                text=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )
            # O processo web nao possui estado necessario para concluir o job.
            # Reinicia-lo enquanto o worker esta vivo prova essa separacao.
            web.terminate()
            web.wait(timeout=10)
            stdout, stderr = worker.communicate(timeout=60)
            self.assertEqual(worker.returncode, 0, stderr or stdout)

            restarted = subprocess.run(
                [sys.executable, "-c", "import app; print('web-restarted')"],
                cwd=BACKEND,
                env=env,
                text=True,
                capture_output=True,
                timeout=60,
                check=False,
            )
            self.assertEqual(restarted.returncode, 0, restarted.stderr)
            with psycopg.connect(DATABASE_URL) as conn:
                status = conn.execute(
                    "SELECT status FROM recalculation_jobs WHERE id=%s", (job_id,)
                ).fetchone()[0]
                self.assertEqual(status, "completed")
                conn.execute("DELETE FROM recalculation_jobs WHERE id=%s", (job_id,))
        finally:
            if web.poll() is None:
                web.terminate()
                try:
                    web.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    web.kill()
            # Pequena espera evita processos filhos ainda fechando sockets no CI.
            time.sleep(0.1)
