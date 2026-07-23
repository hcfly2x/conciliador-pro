import datetime as dt
import io
import sys
import unittest

from openpyxl import load_workbook

from blueprints.system import AUDIT_COLUMNS, _audit_workbook


class AuditWorkbookTests(unittest.TestCase):
    def test_workbook_has_one_typed_row_per_transaction_and_source_links(self):
        row = {key: "" for _, key in AUDIT_COLUMNS}
        row.update(
            transaction_id="tx-1",
            tx_key="key-1",
            date="2026-07-22",
            competence_month="2026-07",
            description="Compra de teste",
            amount=-123.45,
            type="expense",
            status="pending",
            account_name="CARTAO TESTE",
            account_type="credit_card",
            category_name="Alimentação",
            installment_current=2,
            installment_total=3,
            history_match_id="hist-1",
            history_match_confirmed=1,
            imported_file_id="file-1",
            source_filename="fatura.pdf",
            source_file_hash="abc123",
            stored_document_id="doc-1",
            stored_document_filename="fatura.pdf",
        )

        output = _audit_workbook([row], "2026-07-22T12:00:00+00:00")
        workbook = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
        transactions = workbook.worksheets[0]
        headers = [cell.value for cell in transactions[1]]

        self.assertEqual(
            workbook.sheetnames,
            ["Lançamentos", "Resumo", "Legenda", "Dicionário"],
        )
        self.assertEqual(transactions.max_row, 2)
        self.assertEqual(transactions.max_column, 81)
        self.assertIsInstance(transactions.cell(2, headers.index("Data") + 1).value, dt.datetime)
        self.assertEqual(transactions.cell(2, headers.index("Valor") + 1).value, -123.45)
        self.assertEqual(transactions.cell(2, headers.index("Tipo") + 1).value, "Despesa")
        self.assertEqual(
            transactions.cell(2, headers.index("Status") + 1).value,
            "Pendente",
        )
        self.assertEqual(
            transactions.cell(2, headers.index("Tipo de conta") + 1).value,
            "Cartão de crédito",
        )
        self.assertEqual(
            transactions.cell(2, headers.index("Arquivo de origem") + 1).value,
            "fatura.pdf",
        )
        self.assertEqual(
            transactions.cell(2, headers.index("ID vínculo histórico") + 1).value,
            "hist-1",
        )
        technical_column = transactions.column_dimensions[
            transactions.cell(1, headers.index("ID do lançamento") + 1).column_letter
        ]
        self.assertTrue(technical_column.hidden)
        self.assertEqual(transactions.auto_filter.ref, "A1:CC2")
        self.assertEqual(transactions.freeze_panes, "D2")
        self.assertEqual(workbook["Resumo"]["B3"].value, 1)
        self.assertEqual(workbook["Resumo"]["B7"].value, -123.45)
        self.assertEqual(workbook["Legenda"]["A2"].value, "Despesa")
        self.assertEqual(workbook["Dicionário"]["B2"].value, "date")

    def test_workbook_streams_production_volume(self):
        def rows():
            for index in range(5_000):
                row = {key: "" for _, key in AUDIT_COLUMNS}
                row.update(
                    transaction_id=f"tx-{index}",
                    tx_key=f"key-{index}",
                    date="2026-07-22",
                    competence_month="2026-07",
                    description=f"Compra de volume {index}",
                    amount=-10.0,
                    type="expense",
                    status="pending",
                    account_name="CARTAO TESTE",
                    account_type="credit_card",
                )
                yield row

        before_rss = None
        if sys.platform.startswith("linux"):
            import resource

            before_rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss

        output = _audit_workbook(rows(), "2026-07-22T12:00:00+00:00")

        if before_rss is not None:
            import resource

            after_rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
            self.assertLess(after_rss - before_rss, 50 * 1024)

        workbook = load_workbook(
            io.BytesIO(output.getvalue()), data_only=False, read_only=True
        )
        self.assertEqual(
            workbook.sheetnames,
            ["Lançamentos", "Resumo", "Legenda", "Dicionário"],
        )
        self.assertEqual(
            sum(1 for _ in workbook["Lançamentos"].iter_rows()),
            5_001,
        )
        self.assertEqual(workbook["Resumo"]["B3"].value, 5_000)

    def test_workbook_rounds_binary_float_artifacts_for_human_reading(self):
        row = {key: "" for _, key in AUDIT_COLUMNS}
        row.update(
            transaction_id="tx-round",
            amount=-93.98999999999999,
            installment_plan_amount=93.98999999999999,
            type="expense",
            status="pending",
        )

        output = _audit_workbook([row], "2026-07-22T12:00:00+00:00")
        workbook = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
        transactions = workbook["Lançamentos"]
        headers = [cell.value for cell in transactions[1]]

        self.assertEqual(
            transactions.cell(2, headers.index("Valor") + 1).value,
            -93.99,
        )
        self.assertEqual(
            transactions.cell(2, headers.index("Valor do plano") + 1).value,
            93.99,
        )

    def test_workbook_sanitizes_xml_controls_and_timezone_datetimes(self):
        row = {key: "" for _, key in AUDIT_COLUMNS}
        row.update(
            transaction_id="tx-dirty",
            description="Compra\x00 com\x0b controles",
            notes="linha valida\x1flinha invalida",
            amount=-1.0,
            type="expense",
            status="pending",
            imported_at=dt.datetime(
                2026, 7, 22, 9, 30, tzinfo=dt.timezone(dt.timedelta(hours=-3))
            ),
        )

        output = _audit_workbook([row], "2026-07-22T12:00:00+00:00")
        workbook = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
        transactions = workbook["Lançamentos"]
        headers = [cell.value for cell in transactions[1]]

        self.assertEqual(
            transactions.cell(2, headers.index("Descrição") + 1).value,
            "Compra  com  controles",
        )
        self.assertEqual(
            transactions.cell(2, headers.index("Observações") + 1).value,
            "linha valida linha invalida",
        )
        self.assertEqual(
            transactions.cell(2, headers.index("Importado em") + 1).value,
            dt.datetime(2026, 7, 22, 12, 30),
        )


if __name__ == "__main__":
    unittest.main()
